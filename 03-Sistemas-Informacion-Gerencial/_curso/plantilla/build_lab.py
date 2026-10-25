"""Genera un laboratorio de SIG (Prof. Clarence) a partir de la plantilla del profesor.

Uso:
    py -3.13 build_lab.py <spec.json>

El spec (JSON, UTF-8) describe portada y contenido; las rutas de imagen son
relativas a la carpeta del spec. Salida: <salida>.docx y <salida>.pdf en la
carpeta indicada por "salida" (relativa al spec). Word actualiza el índice y
exporta el PDF.

Bloques de "contenido" (uno por elemento de la lista):
    {"h2": "Paso 1: ..."}            subtítulo de paso (entra al índice)
    {"h3": "..."}                     subtítulo menor (entra al índice)
    {"p": "texto"}                    párrafo justificado; **negrita** permitida
    {"nota": "texto"}                 párrafo en cursiva
    {"bullets": ["a", "b"]}           lista con viñetas
    {"pasos": ["a", "b"]}             lista numerada 1., 2., ...
    {"code": "texto"}                 bloque monoespaciado (SQL, fórmulas, hashes)
    {"img": "ruta.png", "ancho": 15, "cap": "Figura 1. ..."}   ancho en cm
    {"tabla": {"titulo": "...", "cols": [...], "filas": [[...], ...]}}
    {"salto": true}                   salto de página

Opcional en la raíz del spec: "tipo" (portada: "LABORATORIO" por defecto, o
"INVESTIGACIÓN", "CUESTIONARIO"...; "lab" es su número); "conclusion" puede
faltar o ir vacía (cuestionarios) y la sección no se crea; "integrantes": ["Apellido Nombre, cédula", ...]
(si falta, se usa INTEGRANTE).
"""
import copy
import json
import re
import sys
from pathlib import Path

import docx
from docx.enum.section import WD_SECTION  # noqa: F401
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
PLANTILLA = AQUI / "PlantillaPresentaciones.docx"
AZUL = RGBColor(0x2F, 0x54, 0x96)

CARRERA = "LICENCIATURA EN DESARROLLO Y GESTIÓN DE SOFTWARE"
ASIGNATURA = "SISTEMAS DE INFORMACIÓN GERENCIAL"
INTEGRANTE = "Gerald Ríos, 8-987-741"


def set_text(p, text, bold=None, size=None):
    """Reemplaza el texto del párrafo conservando el formato del primer run."""
    runs = p.runs
    if not runs:
        r = p.add_run(text)
    else:
        r = runs[0]
        r.text = text
        for extra in runs[1:]:
            extra._element.getparent().remove(extra._element)
    if bold is not None:
        r.bold = bold
    if size is not None:
        r.font.size = Pt(size)
    return r


def estilos(d):
    st = d.styles
    st["Heading 1"].font.bold = True
    st["Heading 1"].font.color.rgb = AZUL
    if "Heading 2" not in [s.name for s in st]:
        h2 = st.add_style("Heading 2", WD_STYLE_TYPE.PARAGRAPH)
    else:
        h2 = st["Heading 2"]
    h2.base_style = st["Heading 1"]
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = AZUL
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True
    ppr = h2.element.get_or_add_pPr()
    for old in ppr.findall(qn("w:outlineLvl")):
        ppr.remove(old)
    lvl = OxmlElement("w:outlineLvl")
    lvl.set(qn("w:val"), "1")
    ppr.append(lvl)

    if "Heading 3" not in [s.name for s in st]:
        h3 = st.add_style("Heading 3", WD_STYLE_TYPE.PARAGRAPH)
    else:
        h3 = st["Heading 3"]
    h3.base_style = h2
    h3.font.size = Pt(12)
    ppr = h3.element.get_or_add_pPr()
    for old in ppr.findall(qn("w:outlineLvl")):
        ppr.remove(old)
    lvl = OxmlElement("w:outlineLvl")
    lvl.set(qn("w:val"), "2")
    ppr.append(lvl)

    # Índice con dos niveles visibles
    for nombre, sangria in (("toc 2", 0.4), ("toc 3", 0.8)):
        if nombre not in [s.name for s in st]:
            s = st.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st["Normal"]
            s.paragraph_format.left_indent = Cm(sangria)
            s.paragraph_format.space_after = Pt(5)


def portada(d, spec):
    ps = d.paragraphs
    lab = spec["lab"]
    integrantes = spec.get("integrantes") or [INTEGRANTE]
    for p in ps[:23]:
        t = p.text.strip()
        if t.startswith("TÉCNICO EN INFORMÁTICA"):
            set_text(p, CARRERA)
        elif t.startswith("SISTEMAS DE INFORMACION"):
            set_text(p, ASIGNATURA)
        elif "Nombre de la Investigación" in t:
            set_text(p, f"{spec.get('tipo', 'LABORATORIO')} #{lab}", bold=True)
            nuevo = copy.deepcopy(p._element)
            p._element.addnext(nuevo)
            q = docx.text.paragraph.Paragraph(nuevo, p._parent)
            set_text(q, f"“{spec['titulo']}”", bold=False)
        elif t.startswith("Integrantes"):
            set_text(p, "Integrantes:" if len(integrantes) > 1 else "Integrante:")
        elif "Su Nombre" in t:
            set_text(p, integrantes[0])
            ancla = p
            for otro in integrantes[1:]:
                nuevo = copy.deepcopy(ancla._element)
                ancla._element.addnext(nuevo)
                ancla = docx.text.paragraph.Paragraph(nuevo, p._parent)
                set_text(ancla, otro)
        elif re.fullmatch(r"\d{4}\.\d{2}\.\d{2}", t):
            set_text(p, spec["fecha"])
    for p in d.paragraphs[23:26]:
        if p.text.strip() == "INDICE":
            r = set_text(p, "INDICE", bold=True, size=16)
            r.font.color.rgb = AZUL


def sin_numero_en_portada(d):
    sec = d.sections[0]
    sec.different_first_page_header_footer = True


def runs_con_negrita(p, texto, italic=False, size=12):
    partes = re.split(r"(\*\*[^*]+\*\*)", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            r = p.add_run(parte[2:-2])
            r.bold = True
        else:
            r = p.add_run(parte)
        r.italic = italic or None
        r.font.size = Pt(size)


class Escritor:
    def __init__(self, d, ancla, base):
        self.d = d
        self.ancla = ancla  # elemento después del cual se inserta
        self.base = base

    def _p(self, style=None):
        p = self.d.add_paragraph(style=style)
        el = p._element
        self.ancla.addnext(el)
        self.ancla = el
        return p

    def _cuerpo(self, p, after=6):
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.15
        pf.space_after = Pt(after)

    def h1(self, texto, salto=True):
        if salto:
            self.salto()
        p = self._p("Heading 1")
        p.add_run(texto)

    def h2(self, texto):
        self._p("Heading 2").add_run(texto)

    def h3(self, texto):
        self._p("Heading 3").add_run(texto)

    def parrafo(self, texto, italic=False, size=12):
        p = self._p()
        self._cuerpo(p)
        runs_con_negrita(p, texto, italic=italic, size=size)

    def lista(self, items, numerada=False):
        for i, it in enumerate(items, 1):
            p = self._p()
            self._cuerpo(p, after=3)
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            runs_con_negrita(p, (f"{i}. " if numerada else "•  ") + it)

    def code(self, texto):
        for linea in texto.split("\n"):
            p = self._p()
            pf = p.paragraph_format
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.5)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F2F2F2")
            p._element.get_or_add_pPr().append(shd)
            r = p.add_run(linea if linea else " ")
            r.font.name = "Consolas"
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Consolas")
            r.font.size = Pt(9.5)
        sep = self._p()
        sep.paragraph_format.space_after = Pt(2)

    def img(self, ruta, ancho, cap):
        ruta = (self.base / ruta).resolve()
        if not ruta.exists():
            raise FileNotFoundError(f"Imagen no encontrada: {ruta}")
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(ruta), width=Cm(ancho))
        if cap:
            c = self._p()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.space_after = Pt(10)
            r = c.add_run(cap)
            r.italic = True
            r.font.size = Pt(10)

    def tabla(self, t):
        if t.get("titulo"):
            self.h3(t["titulo"]) if t.get("en_indice") else self.parrafo(f"**{t['titulo']}**")
        cols, filas = t["cols"], t["filas"]
        tb = self.d.add_table(rows=1 + len(filas), cols=len(cols))
        tb.style = self.d.styles["Table Grid"] if "Table Grid" in [s.name for s in self.d.styles] else None
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, c in enumerate(cols):
            cell = tb.rows[0].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(c))
            r.bold = True
            r.font.size = Pt(10.5)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "D9E2F3")
            cell._element.get_or_add_tcPr().append(shd)
        # una columna con textos largos va toda a la izquierda; las cortas, centradas
        larga = [any(len(str(f[j])) >= 25 for f in filas if j < len(f)) for j in range(len(cols))]
        for i, fila in enumerate(filas, 1):
            for j, v in enumerate(fila):
                cell = tb.rows[i].cells[j]
                cell.text = ""
                r = cell.paragraphs[0].add_run(str(v))
                r.font.size = Pt(10.5)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if larga[j] else WD_ALIGN_PARAGRAPH.CENTER
        # mantener la tabla en una sola página cuando cabe
        for i, row in enumerate(tb.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if i < len(tb.rows) - 1:
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        # bordes aunque falte el estilo Table Grid
        tblPr = tb._element.tblPr
        borders = OxmlElement("w:tblBorders")
        for b in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{b}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "000000")
            borders.append(e)
        tblPr.append(borders)
        self.ancla.addnext(tb._element)
        self.ancla = tb._element
        sep = self._p()
        sep.paragraph_format.space_after = Pt(6)

    def salto(self):
        p = self._p()
        p.add_run().add_break(WD_BREAK.PAGE)


def construir(spec_path):
    spec_path = Path(spec_path).resolve()
    base = spec_path.parent
    spec = json.loads(spec_path.read_text(encoding="utf-8"))

    d = docx.Document(str(PLANTILLA))
    estilos(d)
    portada(d, spec)
    sin_numero_en_portada(d)

    body = d.element.body
    hijos = list(body.iterchildren())
    sdt = next(e for e in hijos if e.tag == qn("w:sdt"))
    # Borra todo lo que hay después del índice (secciones de ejemplo), menos sectPr
    for e in hijos[hijos.index(sdt) + 1:]:
        if e.tag != qn("w:sectPr"):
            body.remove(e)

    w = Escritor(d, sdt, base)
    w.h1("CONTENIDO")
    if spec.get("objetivo"):
        w.parrafo(f"**Objetivo:** {spec['objetivo']}")
    for b in spec["contenido"]:
        if "h2" in b:
            w.h2(b["h2"])
        elif "h3" in b:
            w.h3(b["h3"])
        elif "p" in b:
            w.parrafo(b["p"])
        elif "nota" in b:
            w.parrafo(b["nota"], italic=True, size=11)
        elif "bullets" in b:
            w.lista(b["bullets"])
        elif "pasos" in b:
            w.lista(b["pasos"], numerada=True)
        elif "code" in b:
            w.code(b["code"])
        elif "img" in b:
            w.img(b["img"], b.get("ancho", 15), b.get("cap"))
        elif "tabla" in b:
            w.tabla(b["tabla"])
        elif "salto" in b:
            w.salto()
        else:
            raise ValueError(f"Bloque desconocido: {b}")

    if spec.get("conclusion"):
        w.h1("CONCLUSION")
        for par in spec["conclusion"]:
            w.parrafo(par)

    if spec.get("referencias"):
        w.h1("REFERENCIA")
        for ref in sorted(spec["referencias"], key=lambda s: s.lower()):
            p = w._p()
            p.paragraph_format.left_indent = Cm(1.27)
            p.paragraph_format.first_line_indent = Cm(-1.27)
            p.paragraph_format.space_after = Pt(8)
            runs_con_negrita(p, ref, size=11)

    salida = (base / spec["salida"]).resolve()
    salida.parent.mkdir(parents=True, exist_ok=True)
    docx_path = salida.with_suffix(".docx")
    d.save(str(docx_path))
    entradas = sum(1 for b in spec["contenido"] if "h2" in b or "h3" in b)
    exportar(docx_path, salida.with_suffix(".pdf"), compacto=entradas > 26)
    print(f"OK {docx_path}\nOK {salida.with_suffix('.pdf')}")


def exportar(docx_path, pdf_path, compacto=False):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(docx_path))
        if compacto:  # índice largo (cuestionarios): que quepa en una página
            for nivel in (-20, -21, -22):  # wdStyleTOC1..3
                pf = doc.Styles(nivel).ParagraphFormat
                pf.SpaceAfter = 0
                pf.LineSpacingRule = 0  # wdLineSpaceSingle
        for toc in doc.TablesOfContents:
            toc.Update()
        doc.Fields.Update()
        doc.Save()
        doc.ExportAsFixedFormat(str(pdf_path), 17)  # wdExportFormatPDF
        doc.Close(False)
    finally:
        word.Quit()


if __name__ == "__main__":
    construir(sys.argv[1])
